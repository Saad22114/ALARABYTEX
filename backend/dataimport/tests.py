"""اختبارات الاستيراد من Excel: الترويسات، التحقّق، المعاينة، والتنفيذ."""

from datetime import date
from decimal import Decimal

from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from audit.models import AuditLog
from branches.models import Branch
from core.testsupport import authenticate_admin, make_admin_user
from sale_sessions.models import Employee
from suppliers.models import Fabric, Supplier

from .services import build_rows, build_template, normalize_header, read_sheet
from .specs import ENTITIES


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_bytes(headers, rows, preamble=None):
    """يبني ملف xlsx في الذاكرة مع ترويسة وصفوف (وربما صفوف تمهيدية قبل الترويسة)."""
    workbook = Workbook()
    sheet = workbook.active
    if preamble:
        for line in preamble:
            sheet.append(line)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    import io

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def upload(client, url, buffer, name="data.xlsx"):
    """يرفع الملف كـ ``SimpleUploadedFile`` — مُرمِّز multipart في DRF لا يقرأ
    اسم الملف من الشكل ``(ملف, اسم)`` كما يفعل مُرمِّز Django، فيهمل الاسم."""
    from django.core.files.uploadedfile import SimpleUploadedFile

    payload = SimpleUploadedFile(name, buffer.getvalue(), content_type=XLSX_MIME)
    return client.post(url, {"file": payload}, format="multipart")


class HeaderNormalizationTests(TestCase):
    def test_arabic_variants_are_equivalent(self):
        self.assertEqual(normalize_header("اسم المورد"), normalize_header("اسم المورد"))
        self.assertEqual(normalize_header("  اسم   المورد "), "اسم المورد")
        self.assertEqual(normalize_header("﻿الاسم"), "الاسم")
        self.assertEqual(normalize_header("البريد الإلكتروني"), normalize_header("البريد الالكتروني"))

    def test_english_alias_matches_label(self):
        mapping = ENTITIES["suppliers"].header_map()
        self.assertEqual(mapping[normalize_header("name")].key, "name")
        self.assertEqual(mapping[normalize_header("Name")].key, "name")
        self.assertEqual(mapping[normalize_header("company_name")].key, "company_name")

    def test_all_entities_have_unique_field_keys(self):
        for key, spec in ENTITIES.items():
            keys = [f.key for f in spec.fields]
            self.assertEqual(len(keys), len(set(keys)), key)

    def test_every_field_is_reachable_by_its_own_label(self):
        for key, spec in ENTITIES.items():
            mapping = spec.header_map()
            for field_spec in spec.fields:
                self.assertIn(normalize_header(field_spec.label), mapping, f"{key}.{field_spec.key}")


class BuildRowsTests(TestCase):
    def test_valid_supplier_row_is_create(self):
        rows, _cols, unknown = build_rows("suppliers", ["اسم المورد", "الهاتف"], [["مورد جديد", "0555"]])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].action, "create")
        self.assertEqual(rows[0].errors, [])
        self.assertEqual(rows[0].data["name"], "مورد جديد")
        self.assertEqual(unknown, [])

    def test_existing_supplier_row_is_update(self):
        Supplier.objects.create(name="مورد موجود")
        rows, _c, _u = build_rows("suppliers", ["اسم المورد"], [["مورد موجود"]])
        self.assertEqual(rows[0].action, "update")

    def test_missing_required_field_is_error(self):
        rows, _c, _u = build_rows("suppliers", ["الهاتف"], [["0555"]])
        self.assertEqual(rows[0].action, "skip")
        self.assertTrue(rows[0].errors)
        self.assertIn("اسم المورد", rows[0].errors[0])

    def test_blank_row_produces_error_not_crash(self):
        rows, _c, _u = build_rows("suppliers", ["اسم المورد", "الهاتف"], [[None, None]])
        self.assertEqual(len(rows), 1)
        self.assertFalse(rows[0].valid)

    def test_bad_number_is_reported_with_field_name(self):
        rows, _c, _u = build_rows("fabrics", ["اسم القماش", "سعر الشراء"], [["قطن", "غير رقم"]])
        self.assertFalse(rows[0].valid)
        self.assertIn("سعر الشراء", " ".join(rows[0].errors))

    def test_bad_boolean_is_reported(self):
        rows, _c, _u = build_rows("suppliers", ["اسم المورد", "نشط"], [["مورد", "ربما"]])
        self.assertFalse(rows[0].valid)
        self.assertIn("نشط", " ".join(rows[0].errors))

    def test_choice_accepts_arabic_alias(self):
        rows, _c, _u = build_rows("fabrics", ["اسم القماش", "وحدة البيع"], [["قطن", "لفة"]])
        self.assertEqual(rows[0].data["unit"], "roll")

    def test_choice_rejects_unknown_value(self):
        rows, _c, _u = build_rows("fabrics", ["اسم القماش", "وحدة البيع"], [["قطن", "متر"]])
        self.assertFalse(rows[0].valid)
        self.assertIn("وحدة البيع", " ".join(rows[0].errors))

    def test_unknown_header_is_reported_not_fatal(self):
        rows, cols, unknown = build_rows("suppliers", ["اسم المورد", "عمود غريب"], [["مورد", "x"]])
        self.assertEqual(unknown, ["عمود غريب"])
        self.assertEqual(rows[0].action, "create")

    def test_fabric_matches_by_code_first(self):
        Fabric.objects.create(name="قطن قديم", code="CT-1")
        rows, _c, _u = build_rows("fabrics", ["اسم القماش", "الرمز"], [["قطن جديد", "CT-1"]])
        self.assertEqual(rows[0].action, "update")

    def test_fabric_falls_back_to_name_match(self):
        Fabric.objects.create(name="قطن قديم")
        rows, _c, _u = build_rows("fabrics", ["اسم القماش", "الرمز"], [["قطن قديم", ""]])
        self.assertEqual(rows[0].action, "update")

    def test_unknown_supplier_reference_is_error(self):
        rows, _c, _u = build_rows("fabrics", ["اسم القماش", "المورد"], [["قطن", "مورد غير موجود"]])
        self.assertFalse(rows[0].valid)
        self.assertIn("المورد", " ".join(rows[0].errors))

    def test_known_supplier_reference_is_resolved(self):
        supplier = Supplier.objects.create(name="مورد معروف")
        rows, _c, _u = build_rows("fabrics", ["اسم القماش", "المورد"], [["قطن", "مورد معروف"]])
        self.assertEqual(rows[0].data["supplier"], supplier)

    def test_employee_role_alias_resolved(self):
        rows, _c, _u = build_rows("employees", ["اسم الموظف", "الدور"], [["أحمد", "مبيعات"]])
        self.assertEqual(rows[0].data["role"], "sales")

    def test_date_parsing(self):
        rows, _c, _u = build_rows("employees", ["اسم الموظف", "تاريخ التعيين"], [["أحمد", date(2026, 1, 15)]])
        self.assertEqual(rows[0].data["hire_date"], date(2026, 1, 15))

    def test_date_accepts_iso_string(self):
        rows, _c, _u = build_rows("employees", ["اسم الموظف", "تاريخ التعيين"], [["أحمد", "2026-02-20"]])
        self.assertEqual(rows[0].data["hire_date"], date(2026, 2, 20))

    def test_date_accepts_slash_format(self):
        rows, _c, _u = build_rows("employees", ["اسم الموظف", "تاريخ التعيين"], [["أحمد", "20/02/2026"]])
        self.assertEqual(rows[0].data["hire_date"], date(2026, 2, 20))

    def test_invalid_date_is_reported(self):
        rows, _c, _u = build_rows("employees", ["اسم الموظف", "تاريخ التعيين"], [["أحمد", "ليس تاريخاً"]])
        self.assertFalse(rows[0].valid)
        self.assertIn("تاريخ التعيين", " ".join(rows[0].errors))

    def test_row_index_matches_sheet_row(self):
        rows, _c, _u = build_rows("suppliers", ["اسم المورد"], [["أ"], ["ب"]])
        self.assertEqual([r.index for r in rows], [2, 3])

    def test_numbers_with_commas_are_parsed(self):
        rows, _c, _u = build_rows("suppliers", ["اسم المورد"], [["مورد"]], )
        self.assertTrue(rows[0].valid)
        rows, _c, _u = build_rows("fabrics", ["اسم القماش", "سعر البيع/ياردة"], [["قطن", "1,250.50"]])
        self.assertEqual(rows[0].data["sale_price_yard"], Decimal("1250.50"))

    def test_summary_counts(self):
        Supplier.objects.create(name="موجود")
        rows, _c, _u = build_rows(
            "suppliers", ["اسم المورد"], [["موجود"], ["جديد"], [None]]
        )
        valid = sum(1 for r in rows if r.valid)
        invalid = sum(1 for r in rows if not r.valid)
        self.assertEqual((len(rows), valid, invalid), (3, 2, 1))


class ReadSheetTests(TestCase):
    def test_skips_preamble_rows(self):
        buffer = workbook_bytes(
            ["اسم المورد", "الهاتف"],
            [["مورد", "0555"]],
            preamble=[["تقرير الموردين"], []],
        )
        headers, body = read_sheet(buffer)
        self.assertEqual(headers[0], "اسم المورد")
        self.assertEqual(len(body), 1)

    def test_ignores_blank_rows(self):
        buffer = workbook_bytes(["اسم المورد"], [["مورد"], [None, None], ["آخر"]])
        _headers, body = read_sheet(buffer)
        self.assertEqual(len(body), 2)

    def test_missing_required_header_reports_nothing_valid(self):
        buffer = workbook_bytes(["الهاتف"], [["0555"]])
        headers, body = read_sheet(buffer)
        rows, _c, _u = build_rows("suppliers", headers, body)
        self.assertEqual(len(rows), 1)
        self.assertFalse(rows[0].valid)


class TemplateTests(TestCase):
    def test_template_has_arabic_headers_and_example(self):
        for key in ENTITIES:
            workbook = build_template(key)
            sheet = workbook.active
            headers = [c.value for c in sheet[1] if c.value]
            spec = ENTITIES[key]
            self.assertEqual(headers, [f.label for f in spec.fields], key)
            self.assertIsNotNone(sheet.cell(row=2, column=1).value)

    def test_template_is_readable_back(self):
        import io

        buffer = io.BytesIO()
        build_template("suppliers").save(buffer)
        buffer.seek(0)
        headers, body = read_sheet(buffer)
        rows, _c, unknown = build_rows("suppliers", headers, body)
        self.assertEqual(unknown, [])
        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0].valid, rows[0].errors)


class ImportApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user, self.employee = authenticate_admin(self.client)

    def test_requires_authentication(self):
        self.client.force_authenticate(user=None)
        self.assertEqual(upload(self.client, "/api/import/suppliers/preview/", workbook_bytes(["اسم المورد"], [["م"]])).status_code, 401)

    def test_unknown_entity_is_404(self):
        self.assertEqual(
            self.client.get("/api/import/unicorns/schema/").status_code, 404
        )

    def test_schema_lists_columns(self):
        response = self.client.get("/api/import/suppliers/schema/")
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["entity"], "suppliers")
        self.assertTrue(any(c["required"] for c in payload["columns"]))

    def test_template_downloads_xlsx(self):
        response = self.client.get("/api/import/employees/template/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        workbook = load_workbook(__import__("io").BytesIO(response.content))
        headers = [c.value for c in workbook.active[1] if c.value]
        self.assertIn("اسم الموظف", headers)

    def test_template_requires_view_permission(self):
        perms = dict(self.employee.permissions or {})
        perms["suppliers"] = {"view": False}
        self.employee.permissions = perms
        self.employee.save(update_fields=["permissions"])
        self.assertEqual(self.client.get("/api/import/suppliers/template/").status_code, 403)

    def test_preview_does_not_write(self):
        buffer = workbook_bytes(["اسم المورد", "الهاتف"], [["مورد جديد", "0555"]])
        response = upload(self.client, "/api/import/suppliers/preview/", buffer)
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["summary"]["create"], 1)
        self.assertEqual(payload["summary"]["invalid"], 0)
        self.assertFalse(Supplier.objects.filter(name="مورد جديد").exists())

    def test_commit_creates_rows(self):
        buffer = workbook_bytes(
            ["اسم المورد", "الهاتف", "المدينة"], [["مورد جديد", "0555", "دمشق"]]
        )
        response = upload(self.client, "/api/import/suppliers/commit/", buffer)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["summary"]["created"], 1)
        supplier = Supplier.objects.get(name="مورد جديد")
        self.assertEqual(supplier.city, "دمشق")
        self.assertEqual(supplier.phone, "0555")

    def test_commit_updates_existing_by_name(self):
        Supplier.objects.create(name="مورد", phone="قديم", city="حلب")
        buffer = workbook_bytes(["اسم المورد", "المدينة"], [["مورد", "دمشق"]])
        response = upload(self.client, "/api/import/suppliers/commit/", buffer)
        self.assertEqual(response.json()["summary"]["updated"], 1)
        self.assertEqual(Supplier.objects.count(), 1)
        self.assertEqual(Supplier.objects.get().city, "دمشق")

    def test_invalid_rows_do_not_block_valid_ones(self):
        # الصف الثاني بلا اسم مطلوب، والثالث سعره غير رقمي
        buffer = workbook_bytes(
            ["اسم القماش", "سعر الشراء"],
            [["قطن سليم", 10], [None, 20], ["صوف", "غير رقم"]],
        )
        response = upload(self.client, "/api/import/fabrics/commit/", buffer)
        self.assertEqual(response.status_code, 200)
        summary = response.json()["summary"]
        self.assertEqual(summary["created"], 1)
        # كلا الصفين غير الصالحين رُفضا في التحقّق، فلم يُحاول حفظهما
        self.assertEqual(summary["skipped"], 2)
        self.assertEqual(summary["failed"], 0)
        self.assertTrue(Fabric.objects.filter(name="قطن سليم").exists())
        self.assertFalse(Fabric.objects.filter(name="صوف").exists())

    def test_commit_with_no_valid_rows_returns_400(self):
        buffer = workbook_bytes(["الهاتف"], [["0555"]])
        response = upload(self.client, "/api/import/suppliers/commit/", buffer)
        self.assertEqual(response.status_code, 400)
        self.assertIn("detail", response.json())

    def test_rejects_non_excel_extension(self):
        buffer = workbook_bytes(["اسم المورد"], [["م"]])
        response = upload(self.client, "/api/import/suppliers/preview/", buffer, name="data.csv")
        self.assertEqual(response.status_code, 400)

    def test_missing_file_returns_400(self):
        response = self.client.post("/api/import/suppliers/preview/", {}, format="multipart")
        self.assertEqual(response.status_code, 400)

    def test_import_creates_employee_with_role_permissions(self):
        buffer = workbook_bytes(
            ["اسم الموظف", "الدور", "الراتب الأساسي"],
            [["أحمد", "مبيعات", "500000"]],
        )
        response = upload(self.client, "/api/import/employees/commit/", buffer)
        self.assertEqual(response.status_code, 200)
        employee = Employee.objects.get(name="أحمد")
        self.assertEqual(employee.role, "sales")
        self.assertEqual(employee.base_salary, Decimal("500000"))
        self.assertTrue(employee.has_permission("sales", "view"))
        self.assertFalse(employee.has_permission("payroll", "view"))

    def test_import_links_employee_to_branch(self):
        Branch.objects.create(name="فرع دمشق", code="DM", is_active=False)
        buffer = workbook_bytes(["اسم الموظف", "الفرع"], [["سامي", "فرع دمشق"]])
        response = upload(self.client, "/api/import/employees/commit/", buffer)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Employee.objects.get(name="سامي").branch.name, "فرع دمشق")

    def test_import_records_audit_activity(self):
        buffer = workbook_bytes(["اسم المورد"], [["مورد موثّق"]])
        upload(self.client, "/api/import/suppliers/commit/", buffer)
        log = AuditLog.objects.filter(section="suppliers").order_by("-id").first()
        self.assertIsNotNone(log)
        self.assertIn("استيراد", log.object_repr)
        self.assertEqual(log.changes["details"]["created"], 1)


class ImportPermissionTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user, self.employee = make_admin_user()
        self.client.force_authenticate(user=self.user)

    def _deny(self, section):
        perms = dict(self.employee.permissions or {})
        perms[section] = {"view": True, "create": False, "edit": False, "delete": False}
        self.employee.permissions = perms
        self.employee.save(update_fields=["permissions"])

    def test_cannot_import_without_create_permission(self):
        self._deny("suppliers")
        response = upload(
            self.client, "/api/import/suppliers/commit/", workbook_bytes(["اسم المورد"], [["م"]])
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(Supplier.objects.exists())

    def test_cannot_preview_without_create_permission(self):
        self._deny("fabrics")
        response = upload(
            self.client, "/api/import/fabrics/preview/", workbook_bytes(["اسم القماش"], [["ق"]])
        )
        self.assertEqual(response.status_code, 403)

    def test_can_download_template_with_view_only(self):
        perms = dict(self.employee.permissions or {})
        perms["suppliers"] = {"view": True, "create": False, "edit": False, "delete": False}
        self.employee.permissions = perms
        self.employee.save(update_fields=["permissions"])
        self.assertEqual(self.client.get("/api/import/suppliers/template/").status_code, 200)
